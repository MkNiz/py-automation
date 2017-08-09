import openpyxl

book = openpyxl.load_workbook("dogs.xlsx")
sheet = book.get_sheet_by_name("Dogs")

# Grab data with the reference object
dog_names = openpyxl.chart.Reference(sheet, min_col=1, min_row=2, max_row=8)
dog_ages = openpyxl.chart.Reference(sheet, min_col=3, min_row=2, max_row=8)
# Make a chart and append data to it
chart1 = openpyxl.chart.BarChart()

# Set chart attributes
chart1.title = "Dog Ages"
chart1.type = "col"
chart1.style = 10
chart1.y_axis.title = "Age"
chart1.x_axis.title = "Dog"
chart1.y_axis.scaling.min = 0
chart1.y_axis.scaling.max = 15

chart1.dataLabels = openpyxl.chart.label.DataLabelList() 
chart1.dataLabels.showVal = True

chart1.add_data(dog_ages)
chart1.set_categories(dog_names)
chart1.shape = 4

sheet.add_chart(chart1, "D2")

sheet2 = book.get_sheet_by_name("Adoptions")

# Grab adoption status with a reference object
da_labels = openpyxl.chart.Reference(sheet2, min_col=1, min_row=10, max_row=11)
dog_adopted = openpyxl.chart.Reference(sheet2, min_col=2, min_row=10, max_row=11)
# Make a chart
chart2 = openpyxl.chart.PieChart()

# Chart Attributes
chart2.title = "Dogs Adopted"
chart2.add_data(dog_adopted)
chart2.set_categories(da_labels)

chart2.dataLabels = openpyxl.chart.label.DataLabelList()
chart2.dataLabels.showSerName = True
chart2.dataLabels.showVal = False

sheet2.add_chart(chart2, "D2")

book.save("dogs_charts.xlsx")
print("Saved to dogs_charts.xlsx")